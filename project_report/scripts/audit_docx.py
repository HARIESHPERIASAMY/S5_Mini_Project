"""Comprehensive audit script for Final_Project_Report.docx"""
import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('Final_Project_Report.docx')

# Section 1: Font audit
print("=== FONT/SIZE AUDIT ===")
issues = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    for r in p.runs:
        rt = r.text.strip()
        if not rt:
            continue
        fname = r.font.name
        fsize = r.font.size.pt if r.font.size else None
        # Check if it uses non-TNR font
        if fname and 'Times New Roman' not in fname and fname != 'Courier New':
            issues.append(f'P{i}: non-TNR font "{fname}" in: "{rt[:40]}"')
print(f"Font issues: {len(issues)}")
for iss in issues[:30]:
    print(" ", iss)

# Section 2: Alignment check
print("\n=== ALIGNMENT AUDIT ===")
from docx.enum.text import WD_ALIGN_PARAGRAPH
align_issues = []
body_start = False
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if 'CHAPTER - I' in txt and 'INTRODUCTION' not in txt:
        body_start = True
    if not body_start or not txt:
        continue
    # Skip headings and captions
    if txt.isupper() or txt.startswith('Table') or txt.startswith('Figure') or txt.startswith('['):
        continue
    align = p.alignment
    if align is not None and align != WD_ALIGN_PARAGRAPH.JUSTIFY and align != WD_ALIGN_PARAGRAPH.CENTER:
        align_issues.append(f'P{i}: alignment={align} for: "{txt[:50]}"')

print(f"Alignment issues: {len(align_issues)}")
for iss in align_issues[:20]:
    print(" ", iss)

# Section 3: Chapter structure
print("\n=== CHAPTER STRUCTURE ===")
chapter_headings = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith('CHAPTER') or txt.startswith('REFERENCES') or txt.startswith('APPENDIX'):
        chapter_headings.append(f'P{i}: {txt}')
for h in chapter_headings:
    print(" ", h)

# Section 4: Figure and table cross-reference check
print("\n=== FIGURES AND TABLES ===")
fig_captions = []
table_captions = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith('Figure ') and ':' in txt:
        fig_captions.append(txt)
    elif txt.startswith('Table ') and ':' in txt:
        table_captions.append(txt)

print(f"Figure captions found: {len(fig_captions)}")
for f in fig_captions:
    print(f"  {f[:70]}")
print(f"\nTable captions found: {len(table_captions)}")
for t in table_captions:
    print(f"  {t[:70]}")

# Section 5: References count
print("\n=== REFERENCES ===")
ref_count = 0
in_refs = False
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt == 'REFERENCES':
        in_refs = True
        continue
    if in_refs and txt.startswith('APPENDIX'):
        break
    if in_refs and txt.startswith('['):
        ref_count += 1
        print(f"  {txt[:70]}")
print(f"Total references: {ref_count}")

# Section 6: Abstract word count
print("\n=== ABSTRACT WORD COUNT ===")
in_abstract = False
abstract_text = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt == 'ABSTRACT':
        in_abstract = True
        continue
    if in_abstract and (txt.startswith('TABLE') or txt.startswith('CHAPTER')):
        break
    if in_abstract and txt:
        abstract_text.append(txt)

full_abstract = ' '.join(abstract_text)
word_count = len(full_abstract.split())
print(f"Abstract word count: {word_count} (target: 250-300)")
print("Keywords in abstract:", 'Keywords' in full_abstract or 'keywords' in full_abstract)
