"""
Apply corrections to the newly rebuilt Final_Project_Report.docx (5-chapter structure)
and save to output/SentinelMesh_Final_Report.docx
"""
import sys, shutil, os
sys.stdout.reconfigure(encoding='utf-8')

import docx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Work on the newly rebuilt file
src = 'output/Final_Project_Report.docx'
dst = 'output/SentinelMesh_Final_Report.docx'
shutil.copy(src, dst)
print(f"[1] Copied {src} -> {dst}")

doc = Document(dst)
corrections = []
align_fix = 0
fig_num_fix = 0
table_name_fix = 0

# ---- Fix 1: Paragraph alignment LEFT -> JUSTIFY ----
skip_texts = {
    'CHAPTER - I', 'CHAPTER - II', 'CHAPTER - III', 'CHAPTER - IV', 'CHAPTER - V',
    'INTRODUCTION', 'LITERATURE SURVEY', 'OBJECTIVES AND METHODOLOGY',
    'RESULTS AND DISCUSSION', 'CONCLUSIONS AND SUGGESTIONS FOR FUTURE WORK',
    'BONAFIDE CERTIFICATE', 'DECLARATION', 'ACKNOWLEDGEMENT', 'ABSTRACT',
    'LIST OF FIGURES', 'LIST OF TABLES', 'LIST OF ABBREVIATIONS', 'TABLE OF CONTENTS',
    'REFERENCES', 'BILL OF MATERIALS', 'CODING (CORE IMPLEMENTATION EXCERPTS)',
    'STANDARD TABLES AND SUPPLEMENTARY DATA', 'PUBLICATION CERTIFICATE',
    'PUBLICATION PROOF (MANUSCRIPT FIRST PAGE)', 'WORK CONTRIBUTION',
    'PLAGIARISM REPORT',
}

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt or txt in skip_texts:
        continue
    if txt.startswith('APPENDIX'):
        continue
    if txt.startswith('Figure ') and ':' in txt:
        continue
    if txt.startswith('Table ') and ':' in txt:
        continue
    align = p.alignment
    if align == WD_ALIGN_PARAGRAPH.LEFT:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        align_fix += 1
        corrections.append(f'P{i}: JUSTIFY aligned: "{txt[:50]}"')

print(f"[2] Fixed {align_fix} alignment issues")

# ---- Fix 2: Figure 3.5 -> Figure 3.4 ----
for i, p in enumerate(doc.paragraphs):
    if 'Figure 3.5' in p.text and 'Figure 3.5:' not in p.text:
        for run in p.runs:
            if 'Figure 3.5' in run.text:
                run.text = run.text.replace('Figure 3.5', 'Figure 3.4')
                fig_num_fix += 1
                corrections.append(f'P{i}: Fixed Figure 3.5 -> 3.4')

print(f"[3] Fixed {fig_num_fix} figure number issues")

# ---- Fix 3: Table B.1 -> Table A.1 ----
for i, p in enumerate(doc.paragraphs):
    if 'Table B.1' in p.text:
        for run in p.runs:
            if 'Table B.1' in run.text:
                run.text = run.text.replace('Table B.1', 'Table A.1')
                table_name_fix += 1
                corrections.append(f'P{i}: Fixed Table B.1 -> A.1')

for t_idx, table in enumerate(doc.tables):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if 'Table B.1' in p.text:
                    for run in p.runs:
                        if 'Table B.1' in run.text:
                            run.text = run.text.replace('Table B.1', 'Table A.1')
                            table_name_fix += 1

print(f"[4] Fixed {table_name_fix} table naming issues")

# ---- Save ----
doc.save(dst)
print(f"[5] Saved corrected document -> {dst}")
print(f"\nTotal corrections: {len(corrections)}")
print(f"  - Alignment: {align_fix}")
print(f"  - Figure refs: {fig_num_fix}")
print(f"  - Table names: {table_name_fix}")
