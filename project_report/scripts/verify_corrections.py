import docx, sys
sys.stdout.reconfigure(encoding='utf-8')
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document('output/SentinelMesh_Final_Report.docx')
print('=== VERIFICATION: Alignment check after correction ===')
issues = []
skip_texts = {
    'CHAPTER - I', 'CHAPTER - II', 'CHAPTER - III', 'CHAPTER - IV', 
    'CHAPTER - V', 'CHAPTER - VI', 'REFERENCES', 'BONAFIDE CERTIFICATE',
    'DECLARATION', 'ACKNOWLEDGEMENT', 'ABSTRACT', 'LIST OF FIGURES',
    'LIST OF TABLES', 'LIST OF ABBREVIATIONS', 'TABLE OF CONTENTS',
    'INTRODUCTION', 'LITERATURE SURVEY', 'OBJECTIVES AND METHODOLOGY',
    'PROPOSED WORK MODULES', 'RESULTS AND DISCUSSION',
    'CONCLUSIONS AND SUGGESTIONS FOR FUTURE WORK'
}
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt or txt in skip_texts:
        continue
    if txt.startswith('Figure ') and ':' in txt:
        continue
    if txt.startswith('Table ') and ':' in txt:
        continue
    align = p.alignment
    if align == WD_ALIGN_PARAGRAPH.LEFT:
        issues.append(f'P{i}: LEFT: "{txt[:60]}"')

print(f'Remaining alignment issues: {len(issues)}')
for iss in issues[:15]:
    print(' ', iss)

print()
print('=== VERIFICATION: Figure 3.5 references ===')
fig35_found = []
for i, p in enumerate(doc.paragraphs):
    if 'Figure 3.5' in p.text:
        fig35_found.append(f'P{i}: {p.text[:70]}')
print(f'Figure 3.5 references remaining: {len(fig35_found)}')
for f in fig35_found:
    print(' ', f)

print()
print('=== VERIFICATION: Table B.1 references ===')
tableb1_found = []
for i, p in enumerate(doc.paragraphs):
    if 'Table B.1' in p.text:
        tableb1_found.append(f'P{i}: {p.text[:70]}')
print(f'Table B.1 references remaining: {len(tableb1_found)}')
for t in tableb1_found:
    print(' ', t)

print()
print('Verification complete. Document: output/SentinelMesh_Final_Report.docx')
